from io import BytesIO
from pathlib import Path

import pymupdf
from docx import Document


def extract_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()

    if extension == ".pdf":
        with pymupdf.open(
            stream=content,
            filetype="pdf"
        ) as document:
            text = "\n".join(
                page.get_text() for page in document
            )

    elif extension == ".docx":
        document = Document(BytesIO(content))

        parts = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                parts.append(
                    " | ".join(
                        cell.text.strip()
                        for cell in row.cells
                    )
                )

        text = "\n".join(parts)

    else:
        raise ValueError(
            "Only PDF and DOCX files are supported."
        )

    if not text.strip():
        raise ValueError(
            "No readable text found in the document. "
            "The PDF may be scanned and require OCR."
        )

    return text.strip()